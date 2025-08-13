import os
import psycopg2
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from io import BytesIO
import pandas as pd
import numpy as np
from datetime import datetime
from dateutil import parser
from flask_caching import Cache
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment

app = Flask(__name__)
app.secret_key = os.urandom(24)
cache = Cache(app, config={'CACHE_TYPE': 'SimpleCache'})
BASE_DIR = "amc"
required_cols = {'Name', 'SectorName', 'NoOfShare', 'Month', 'SharesZG', 'MarketValue', 'MarketValueZG', 'HoldingPercentage'}

def get_cache_key(min_shares):
    """Generate a cache key based on file modification times and min_shares."""
    latest_mtime = 0
    for amc_dir in os.listdir(BASE_DIR):
        amc_path = os.path.join(BASE_DIR, amc_dir)
        if os.path.isdir(amc_path):
            for scheme_file in os.listdir(amc_path):
                if scheme_file.endswith('.csv'):
                    file_path = os.path.join(amc_path, scheme_file)
                    mtime = os.path.getmtime(file_path)
                    latest_mtime = max(latest_mtime, mtime)
    return f"all_data_{min_shares}_{int(latest_mtime)}"

@cache.memoize(timeout=3600)  # Cache for 1 hour
def load_all_data(min_shares, cache_key):
    """Load and concatenate all CSV files, using cache_key to ensure freshness."""
    all_dfs = []
    dtypes = {
        'NoOfShare': 'int32',
        'SharesZG': 'float32',
        'MarketValue': 'float32',
        'MarketValueZG': 'float32',
        'HoldingPercentage': 'float32'
    }
    for amc_dir in os.listdir(BASE_DIR):
        amc_path = os.path.join(BASE_DIR, amc_dir)
        if os.path.isdir(amc_path):
            for scheme_file in os.listdir(amc_path):
                if scheme_file.endswith('.csv'):
                    file_path = os.path.join(amc_path, scheme_file)
                    try:
                        df = pd.read_csv(file_path, usecols=required_cols, dtype=dtypes)
                        if required_cols.issubset(df.columns):
                            df = df[df['NoOfShare'] >= min_shares]
                            df['AMC'] = amc_dir
                            df['Scheme'] = scheme_file
                            all_dfs.append(df)
                    except Exception as e:
                        print(f"Error reading {file_path}: {e}")
    return pd.concat(all_dfs, ignore_index=True) if all_dfs else pd.DataFrame()

def init_db():
    try:
        conn = psycopg2.connect(os.environ.get('DATABASE_URL'))
        with conn.cursor() as cur:
            cur.execute('''
                CREATE TABLE IF NOT EXISTS user_logins (
                    email TEXT NOT NULL,
                    timestamp TEXT NOT NULL
                )
            ''')
            conn.commit()
        conn.close()
    except Exception as e:
        print(f"Error initializing database: {e}")

def store_email(email):
    try:
        conn = psycopg2.connect(os.environ.get('DATABASE_URL'))
        with conn.cursor() as cur:
            timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            cur.execute('INSERT INTO user_logins (email, timestamp) VALUES (%s, %s)', (email, timestamp))
            conn.commit()
        conn.close()
    except Exception as e:
        print(f"Error storing email: {e}")

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.json.get('email')
        if email and '@' in email and '.' in email:
            session['user_email'] = email
            store_email(email)
            return jsonify({"success": True, "redirect": url_for('home')})
        return jsonify({"success": False, "message": "Invalid email"})
    return render_template("login.html")

# Initialize database when app starts
init_db()

@app.route('/')
def home():
    if 'user_email' not in session:
        return redirect(url_for('login'))
    return render_template("index.html")

@app.route('/get_amcs', methods=['GET'])
def get_amcs():
    if 'user_email' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    try:
        amcs = sorted([f for f in os.listdir(BASE_DIR) if os.path.isdir(os.path.join(BASE_DIR, f))])
    except FileNotFoundError:
        amcs = []
    return jsonify(amcs)

@app.route('/get_schemes', methods=['POST'])
def get_schemes():
    if 'user_email' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    amc = request.json.get('amc')
    amc_path = os.path.join(BASE_DIR, amc)
    if not os.path.exists(amc_path):
        return jsonify([])
    
    schemes = [f for f in os.listdir(amc_path) if f.endswith('.csv')]
    return jsonify(schemes)

@app.route('/get_data_for_share_funds', methods=['POST'])
def get_data_for_share_funds():
    if 'user_email' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    share = request.json.get('share')
    min_shares = 1
    try:
        cache_key = get_cache_key(min_shares)
        df = load_all_data(min_shares, cache_key)
        if df.empty:
            return jsonify({"html": "<div class='text-danger'>No valid CSV files found</div>"})

        df = df[df['Name'] == share]
        if df.empty:
            return jsonify({"html": "<div class='text-danger'>No funds found for this share</div>"})

        # Create pivot table for number of shares
        pivot = df.pivot_table(
            index=['AMC', 'Scheme'],
            columns='Month',
            values='NoOfShare',
            aggfunc="sum",
            fill_value=0
        )
        pivot.columns = [str(col).strip() for col in pivot.columns]
        
        # Sort columns by date
        def try_parse(col):
            try:
                dt = parser.parse(col, dayfirst=False, fuzzy=True)
                return (col, dt)
            except:
                return None

        parsed = list(filter(None, map(try_parse, pivot.columns)))
        sorted_cols = [col for col, _ in sorted(parsed, key=lambda x: x[1])]
        pivot = pivot[sorted_cols]
        pivot.reset_index(inplace=True)
        pivot.columns.name = None

        # Generate HTML table
        green_shades = [
            "#e9fbe9", "#c8f7c5", "#a3f3a3", "#6de26d", "#36c836", "#1e9f1e", "#107a10", "#075c07", "#033b03"
        ]
        red_shades = [
            "#ffe6e6", "#ffc2c2", "#ff9999", "#ff6b6b", "#ff3b3b", "#e60000", "#990000", "#660000", "#330000"
        ]

        def get_trend_colors(values):
            colors = [green_shades[0]] * len(values)
            trend_count = 1
            direction = 'up'
            for i in range(1, len(values)):
                if pd.notnull(values[i]) and pd.notnull(values[i-1]):
                    if values[i] > values[i-1]:
                        if direction == 'up':
                            trend_count += 1
                        else:
                            trend_count = 1
                            direction = 'up'
                        colors[i] = green_shades[min(trend_count, len(green_shades)-1)]
                    elif values[i] < values[i-1]:
                        if direction == 'down':
                            trend_count += 1
                        else:
                            trend_count = 1
                            direction = 'down'
                        colors[i] = red_shades[min(trend_count, len(red_shades)-1)]
                    else:
                        colors[i] = colors[i-1]
            return colors

        html = f"<h5>Funds Holding {share} (Shares in Lakhs)</h5>"
        html += "<table class='table table-bordered table-striped'><thead><tr>"
        columns = ['AMC', 'Scheme'] + sorted_cols
        for col in columns:
            html += f"<th>{col}</th>"
        html += "</tr></thead><tbody>"

        for _, row in pivot.iterrows():
            html += "<tr>"
            html += f"<td>{row['AMC']}</td><td>{row['Scheme']}</td>"
            values = [row[month] / 100000 for month in sorted_cols]  # Convert to lakhs
            colors = get_trend_colors(values)
            for value, color in zip(values, colors):
                html += f"<td style='background-color:{color}'>{value:.2f} L</td>"
            html += "</tr>"

        html += "</tbody></table>"
        return jsonify({"html": html})
    except Exception as e:
        return jsonify({"html": f"<div class='text-danger'>Error loading funds data: {e}</div>"})

@app.route('/get_funds_for_share', methods=['POST'])
def get_funds_for_share():
    if 'user_email' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    share = request.json.get('share')
    min_shares = 1
    try:
        cache_key = get_cache_key(min_shares)
        df = load_all_data(min_shares, cache_key)
        if df.empty:
            return jsonify([])
        funds = df[df['Name'] == share][['AMC', 'Scheme']].drop_duplicates()
        fund_names = funds.apply(lambda x: f"{x['AMC']} - {x['Scheme']}", axis=1).tolist()
        return jsonify(fund_names)
    except Exception as e:
        return jsonify([])

@app.route('/get_data', methods=['POST'])
def get_data():
    if 'user_email' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    amc = request.json.get('amc')
    scheme = request.json.get('scheme')
    min_shares = int(request.json.get('min_shares', 0))
    view = request.json.get('view', 'holding_percentage')

    dtypes = {
        'NoOfShare': 'int32',
        'SharesZG': 'float32',
        'MarketValue': 'float32',
        'MarketValueZG': 'float32',
        'HoldingPercentage': 'float32'
    }

    if view in ['sector_holding_all', 'share_wise_shares_all']:
        try:
            cache_key = get_cache_key(min_shares)
            df = load_all_data(min_shares, cache_key)
            if df.empty:
                return jsonify({"html": "<div class='text-danger'>No valid CSV files found</div>"})
        except Exception as e:
            return jsonify({"html": f"<div class='text-danger'>Error loading CSVs: {e}</div>"})
    else:
        file_path = os.path.join(BASE_DIR, amc, scheme)
        try:
            df = pd.read_csv(file_path, usecols=required_cols, dtype=dtypes)
        except Exception as e:
            return jsonify({"html": f"<div class='text-danger'>Error loading CSV: {e}</div>"})
        if not required_cols.issubset(df.columns):
            return jsonify({"html": "<div class='text-danger'>Invalid CSV format: Missing required columns</div>"})
        df = df[df['NoOfShare'] >= min_shares]

    def create_pivot_table(value_col, index_cols=['Name', 'SectorName']):
        pivot = df.pivot_table(
            index=index_cols,
            columns='Month',
            values=value_col,
            aggfunc="sum",
            fill_value=0
        )
        pivot.columns = [str(col).strip() for col in pivot.columns]
        
        def try_parse(col):
            try:
                dt = parser.parse(col, dayfirst=False, fuzzy=True)
                return (col, dt)
            except:
                return None

        parsed = list(filter(None, map(try_parse, pivot.columns)))
        sorted_cols = [col for col, _ in sorted(parsed, key=lambda x: x[1])]
        pivot = pivot[sorted_cols]
        pivot.reset_index(inplace=True)
        if len(index_cols) == 2:
            pivot.rename(columns={"Name": "Share", "SectorName": "Sector"}, inplace=True)
        else:
            pivot.rename(columns={"SectorName": "Sector"}, inplace=True)
        pivot.columns.name = None
        return pivot, sorted_cols

    # Compute stats boxes if applicable
    boxes = []
    if view in ['holding_percentage', 'sector_holding', 'consolidated']:
        months = sorted(set(df['Month']), key=lambda x: parser.parse(x))
        if months:
            latest_month = months[-1]
            df_latest = df[df['Month'] == latest_month]
            # Box 1: Highest Allocation
            top_alloc = df_latest.sort_values('HoldingPercentage', ascending=False).head(5)
            items = [f"{i+1}. {row['Name']}" for i, (_, row) in enumerate(top_alloc.iterrows())]
            box_html = f'<div class="card m-1 stats-box" style="flex: 1; background-color: #e3f2fd;"><div class="card-body"><h5 class="card-title">Highest Allocation</h5><ul class="list-group list-group-flush">' + ''.join(f'<li class="list-group-item">{item}</li>' for item in items) + '</ul></div></div>'
            boxes.append(box_html)

            if len(months) >= 2:
                prev_month = months[-2]
                df_prev = df[df['Month'] == prev_month]

                # Box 2: New Entry
                new_shares = set(df_latest['Name']) - set(df_prev['Name'])
                new_df = df_latest[df_latest['Name'].isin(new_shares)].sort_values('HoldingPercentage', ascending=False).head(5)
                items = [f"{i+1}. {row['Name']}" for i, (_, row) in enumerate(new_df.iterrows())]
                box_html = f'<div class="card m-1 stats-box" style="flex: 1; background-color: #e8f5e9;"><div class="card-body"><h5 class="card-title">New Entry</h5><ul class="list-group list-group-flush">' + ''.join(f'<li class="list-group-item">{item}</li>' for item in items) + '</ul></div></div>'
                boxes.append(box_html)

                # Box 3: Completely Exited
                exited_shares = set(df_prev['Name']) - set(df_latest['Name'])
                exited_df = df_prev[df_prev['Name'].isin(exited_shares)].sort_values('HoldingPercentage', ascending=False).head(5)
                items = [f"{i+1}. {row['Name']}" for i, (_, row) in enumerate(exited_df.iterrows())]
                box_html = f'<div class="card m-1 stats-box" style="flex: 1; background-color: #ffebee;"><div class="card-body"><h5 class="card-title">Completely Exited</h5><ul class="list-group list-group-flush">' + ''.join(f'<li class="list-group-item">{item}</li>' for item in items) + '</ul></div></div>'
                boxes.append(box_html)

                # Box 4: Increasing Stake
                merged = pd.merge(df_latest[['Name', 'HoldingPercentage']], df_prev[['Name', 'HoldingPercentage']], on='Name', how='inner', suffixes=('_latest', '_prev'))
                merged['change'] = merged['HoldingPercentage_latest'] - merged['HoldingPercentage_prev']
                increases = merged[merged['change'] > 0].sort_values('change', ascending=False).head(5)
                items = [f"{i+1}. {row['Name']}" for i, (_, row) in enumerate(increases.iterrows())]
                box_html = f'<div class="card m-1 stats-box" style="flex: 1; background-color: #f1f8e9;"><div class="card-body"><h5 class="card-title">Increasing Stake</h5><ul class="list-group list-group-flush">' + ''.join(f'<li class="list-group-item">{item}</li>' for item in items) + '</ul></div></div>'
                boxes.append(box_html)

                # Box 5: Decreasing Stake
                decreases = merged[merged['change'] < 0].sort_values('change', ascending=True).head(5)
                items = [f"{i+1}. {row['Name']}" for i, (_, row) in enumerate(decreases.iterrows())]
                box_html = f'<div class="card m-1 stats-box" style="flex: 1; background-color: #fff3e0;"><div class="card-body"><h5 class="card-title">Decreasing Stake</h5><ul class="list-group list-group-flush">' + ''.join(f'<li class="list-group-item">{item}</li>' for item in items) + '</ul></div></div>'
                boxes.append(box_html)

    # Create only the necessary pivot table
    if view == 'holding_percentage':
        pivot, month_cols = create_pivot_table('HoldingPercentage')
    elif view == 'sector_holding':
        pivot, month_cols = create_pivot_table('HoldingPercentage', index_cols=['SectorName'])
    elif view == 'sector_holding_all':
        pivot, month_cols = create_pivot_table('HoldingPercentage', index_cols=['SectorName'])
    elif view == 'share_wise_shares_all':
        pivot, month_cols = create_pivot_table('NoOfShare', index_cols=['Name', 'SectorName'])
    elif view == 'consolidated':
        pivot_no_shares, month_cols = create_pivot_table('NoOfShare')
        pivot_shares_zg, _ = create_pivot_table('SharesZG')
        pivot_market_value, _ = create_pivot_table('MarketValue')
        pivot_market_value_zg, _ = create_pivot_table('MarketValueZG')
        pivot_holding_percentage, _ = create_pivot_table('HoldingPercentage')
    else:
        return jsonify({"html": "<div class='text-danger'>Invalid view selected</div>"})

    green_shades = [
        "#e9fbe9", "#c8f7c5", "#a3f3a3", "#6de26d", "#36c836", "#1e9f1e", "#107a10", "#075c07", "#033b03"
    ]
    red_shades = [
        "#ffe6e6", "#ffc2c2", "#ff9999", "#ff6b6b", "#ff3b3b", "#e60000", "#990000", "#660000", "#330000"
    ]

    def get_trend_colors(values):
        colors = [green_shades[0]] * len(values)
        trend_count = 1
        direction = 'up'
        for i in range(1, len(values)):
            if pd.notnull(values[i]) and pd.notnull(values[i-1]):
                if values[i] > values[i-1]:
                    if direction == 'up':
                        trend_count += 1
                    else:
                        trend_count = 1
                        direction = 'up'
                    colors[i] = green_shades[min(trend_count, len(green_shades)-1)]
                elif values[i] < values[i-1]:
                    if direction == 'down':
                        trend_count += 1
                    else:
                        trend_count = 1
                        direction = 'down'
                    colors[i] = red_shades[min(trend_count, len(red_shades)-1)]
                else:
                    colors[i] = colors[i-1]
        return colors

    def generate_table(pivot, title, is_decimal=False, is_consolidated=False, pivots=None):
        static_cols = ["Share", "Sector"] if view not in ['sector_holding', 'sector_holding_all'] else ["Sector"]
        columns = static_cols + month_cols if not is_consolidated else ["Share", "Sector", "Metric"] + month_cols
        html = f"<h4>{title}</h4>"
        html += "<div style='display:block; overflow-x:auto; width:100%'><table class='table table-bordered table-striped'><thead><tr>"

        for col in columns:
            html += f"<th>{col}</th>"
        html += "</tr></thead><tbody>"

        if is_consolidated:
            metrics = [
                ("No. of Shares (in L)", pivot_no_shares, False),
                ("Changes in Holding %", pivot_shares_zg, True),
                ("Market Value", pivot_market_value, True),
                ("Changes in Market Value %", pivot_market_value_zg, True),
                ("% of Total Holding", pivot_holding_percentage, True)
            ]

            grouped = pivot_no_shares.groupby(['Share', 'Sector'])
            for (share, sector), group in grouped:
                rowspan = len(metrics)
                for metric_idx, (metric_name, pivot_data, is_decimal) in enumerate(metrics):
                    metric_row = pivot_data[(pivot_data['Share'] == share) & (pivot_data['Sector'] == sector)]
                    html += "<tr>"
                    if metric_idx == 0:
                        html += f"<td rowspan='{rowspan}' style='text-align:center; vertical-align:middle;'>{share}</td>"
                        html += f"<td rowspan='{rowspan}' style='text-align:center; vertical-align:middle;'>{sector}</td>"

                    html += f"<td>{metric_name}</td>"

                    values = [metric_row[month].iloc[0] if not metric_row.empty else 0.0 for month in month_cols]
                    if metric_name == "No. of Shares (in L)":
                        values = [value / 100000 for value in values]  # Convert to lakhs
                    colors = get_trend_colors(values)
                    for value, color in zip(values, colors):
                        display_value = f"{value:.2f}" if is_decimal or metric_name == "No. of Shares (in L)" else str(int(value))
                        if metric_name == "No. of Shares (in L)":
                            display_value += " L"
                        html += f"<td style='background-color:{color}'>{display_value}</td>"
                    html += "</tr>"

        else:
            for _, row in pivot.iterrows():
                html += "<tr>"
                if view in ['sector_holding', 'sector_holding_all']:
                    html += f"<td>{row['Sector']}</td>"
                else:
                    share_cell_class = "share-cell" if view == 'share_wise_shares_all' else ""
                    html += f"<td class='{share_cell_class}'>{row['Share']}</td><td>{row['Sector']}</td>"

                values = [row[month] for month in month_cols]
                if view == 'share_wise_shares_all':
                    values = [value / 100000 for value in values]  # Convert to lakhs
                colors = get_trend_colors(values)
                for value, color in zip(values, colors):
                    display_value = f"{value:.2f} L" if view == 'share_wise_shares_all' else f"{value:.2f}" if is_decimal else str(int(value))
                    html += f"<td style='background-color:{color}'>{display_value}</td>"
                html += "</tr>"

        html += "</tbody></table></div>"
        return html

    if view == 'holding_percentage':
        html = generate_table(pivot, "% of Total Holding", is_decimal=True)
    elif view == 'sector_holding':
        html = generate_table(pivot, "Sector Wise Holding %", is_decimal=True)
    elif view == 'sector_holding_all':
        html = generate_table(pivot, "Sector Wise Holding % (All AMCs)", is_decimal=True)
    elif view == 'share_wise_shares_all':
        html = generate_table(pivot, "Share Wise Number of Shares (All AMCs)")
    elif view == 'consolidated':
        html = generate_table(pivot_no_shares, "Consolidated View", is_consolidated=True, 
                             pivots=[pivot_no_shares, pivot_shares_zg, pivot_market_value, 
                                     pivot_market_value_zg, pivot_holding_percentage])
    else:
        html = "<div class='text-danger'>Invalid view selected</div>"

    return jsonify({"html": html, "boxes": boxes})

@app.route('/export_to_excel', methods=['POST'])
def export_to_excel():
    if 'user_email' not in session:
        return jsonify({"error": "Unauthorized"}), 401
    
    amc = request.json.get('amc')
    scheme = request.json.get('scheme')
    min_shares = int(request.json.get('min_shares', 0))
    view = request.json.get('view', 'holding_percentage')

    dtypes = {
        'NoOfShare': 'int32',
        'SharesZG': 'float32',
        'MarketValue': 'float32',
        'MarketValueZG': 'float32',
        'HoldingPercentage': 'float32'
    }

    # Load data based on view
    if view in ['sector_holding_all', 'share_wise_shares_all']:
        try:
            cache_key = get_cache_key(min_shares)
            df = load_all_data(min_shares, cache_key)
            if df.empty:
                return jsonify({"error": "No valid CSV files found"}), 400
        except Exception as e:
            return jsonify({"error": f"Error loading CSVs: {e}"}), 500
    else:
        file_path = os.path.join(BASE_DIR, amc, scheme)
        try:
            df = pd.read_csv(file_path, usecols=required_cols, dtype=dtypes)
        except Exception as e:
            return jsonify({"error": f"Error loading CSV: {e}"}), 500
        if not required_cols.issubset(df.columns):
            return jsonify({"error": "Invalid CSV format: Missing required columns"}), 400
        df = df[df['NoOfShare'] >= min_shares]

    def create_pivot_table(value_col, index_cols=['Name', 'SectorName']):
        pivot = df.pivot_table(
            index=index_cols,
            columns='Month',
            values=value_col,
            aggfunc="sum",
            fill_value=0
        )
        pivot.columns = [str(col).strip() for col in pivot.columns]
        
        def try_parse(col):
            try:
                dt = parser.parse(col, dayfirst=False, fuzzy=True)
                return (col, dt)
            except:
                return None

        parsed = list(filter(None, map(try_parse, pivot.columns)))
        sorted_cols = [col for col, _ in sorted(parsed, key=lambda x: x[1])]
        pivot = pivot[sorted_cols]
        pivot.reset_index(inplace=True)
        if len(index_cols) == 2:
            pivot.rename(columns={"Name": "Share", "SectorName": "Sector"}, inplace=True)
        else:
            pivot.rename(columns={"SectorName": "Sector"}, inplace=True)
        pivot.columns.name = None
        return pivot, sorted_cols

    try:
        # Create pivot table based on view
        if view == 'holding_percentage':
            pivot, month_cols = create_pivot_table('HoldingPercentage')
        elif view == 'sector_holding':
            pivot, month_cols = create_pivot_table('HoldingPercentage', index_cols=['SectorName'])
        elif view == 'sector_holding_all':
            pivot, month_cols = create_pivot_table('HoldingPercentage', index_cols=['SectorName'])
        elif view == 'share_wise_shares_all':
            pivot, month_cols = create_pivot_table('NoOfShare', index_cols=['Name', 'SectorName'])
            pivot[month_cols] = pivot[month_cols] / 100000  # Convert to lakhs
        elif view == 'consolidated':
            # Create pivot tables for each metric
            pivot_no_shares, month_cols = create_pivot_table('NoOfShare')
            pivot_no_shares[month_cols] = pivot_no_shares[month_cols] / 100000  # Convert to lakhs
            pivot_shares_zg, _ = create_pivot_table('SharesZG')
            pivot_market_value, _ = create_pivot_table('MarketValue')
            pivot_market_value_zg, _ = create_pivot_table('MarketValueZG')
            pivot_holding_percentage, _ = create_pivot_table('HoldingPercentage')
            
            # Combine all metrics into a single DataFrame, grouping by Share
            combined_dfs = []
            metrics = [
                ("No. of Shares (in L)", pivot_no_shares, False),
                ("Changes in Holding %", pivot_shares_zg, True),
                ("Market Value", pivot_market_value, True),
                ("Changes in Market Value %", pivot_market_value_zg, True),
                ("% of Total Holding", pivot_holding_percentage, True)
            ]
            for metric_name, pivot, is_decimal in metrics:
                pivot_copy = pivot.copy()
                pivot_copy['Metric'] = metric_name
                # Reorder columns to match UI: Share, Sector, Metric, then months
                pivot_copy = pivot_copy[['Share', 'Sector', 'Metric'] + month_cols]
                if is_decimal:
                    pivot_copy[month_cols] = pivot_copy[month_cols].round(2)
                combined_dfs.append(pivot_copy)
            
            # Concatenate and sort by Share, then Metric to group all metrics for each share
            combined_df = pd.concat(combined_dfs, ignore_index=True)
            combined_df = combined_df.sort_values(by=['Share', 'Metric'])
        else:
            return jsonify({"error": "Invalid view selected"}), 400

        # Create Excel file
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            if view == 'consolidated':
                combined_df.to_excel(writer, sheet_name='Consolidated View', index=False)
                worksheet = writer.sheets['Consolidated View']
                
                # Format cells: Append "L" to No. of Shares and align numbers
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=4):
                    for cell in row:
                        if cell.column >= 4:  # Month columns
                            value = cell.value
                            if value is not None and not np.isnan(value):
                                if worksheet.cell(row=cell.row, column=3).value == 'No. of Shares (in L)':
                                    cell.value = f"{value:.2f} L"
                                    cell.alignment = Alignment(horizontal='right')
                                else:
                                    cell.value = f"{value:.2f}"
                                    cell.alignment = Alignment(horizontal='right')
                
                # Adjust column widths
                for col in worksheet.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = max_length + 2
                    worksheet.column_dimensions[column].width = adjusted_width
            else:
                pivot.to_excel(writer, sheet_name=view, index=False)
                worksheet = writer.sheets[view]
                
                # Format numeric columns
                for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=3):
                    for cell in row:
                        if cell.column >= 3:  # Month columns
                            value = cell.value
                            if value is not None and not np.isnan(value):
                                if view == 'share_wise_shares_all':
                                    cell.value = f"{value:.2f} L"
                                    cell.alignment = Alignment(horizontal='right')
                                else:
                                    cell.value = f"{value:.2f}"
                                    cell.alignment = Alignment(horizontal='right')
                
                # Adjust column widths
                for col in worksheet.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = max_length + 2
                    worksheet.column_dimensions[column].width = adjusted_width

        output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'table_data_{view}.xlsx'
        )

    except Exception as e:
        return jsonify({"error": f"Error generating Excel: {e}"}), 500

if __name__ == '__main__':
    app.run(debug=True)